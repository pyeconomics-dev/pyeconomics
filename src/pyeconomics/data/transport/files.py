# src/pyeconomics/data/transport/files.py
# Copyright 2026 Nathan Ramos, CFA
# SPDX-License-Identifier: Apache-2.0
"""Bounded in-memory file readers; never extract an archive to the filesystem."""

from __future__ import annotations

from io import BytesIO
from typing import Any
from zipfile import ZipFile

from pyeconomics.data.errors import ResponseError
from pyeconomics.data.redaction import call_scrubbed

MAX_EXPANDED_BYTES = 96 * 1024 * 1024
MAX_ARCHIVE_ENTRIES = 256
MAX_SHEET_ROWS = 100_100
MAX_SHEET_COLUMNS = 1000


def archive(body: bytes, provider: str) -> dict[str, bytes]:
    """Read a bounded archive in memory; reject encrypted or oversized members."""

    def read() -> dict[str, bytes]:
        with ZipFile(BytesIO(body)) as zipped:
            members = zipped.infolist()
            if (
                len(members) > MAX_ARCHIVE_ENTRIES
                or sum(m.file_size for m in members) > MAX_EXPANDED_BYTES
            ):
                msg = "archive exceeds the expanded size or entry cap"
                raise ValueError(msg)
            if any(m.flag_bits & 1 for m in members):
                msg = "encrypted archives are unsupported"
                raise ValueError(msg)
            return {m.filename: zipped.read(m) for m in members}

    return call_scrubbed(
        read, lambda message: ResponseError(message, provider=provider)
    )


def spreadsheet(body: bytes, sheet: str, provider: str) -> tuple[tuple[Any, ...], ...]:
    """Load an XLSX worksheet lazily, read-only, with expanded-size bounds."""
    # Validate the compressed container before openpyxl touches XML. openpyxl
    # discovers defusedxml at import; no workbook is written to disk.
    archive(body, provider)

    def read() -> tuple[tuple[Any, ...], ...]:
        from openpyxl import load_workbook  # noqa: PLC0415 - optional browser import

        book = load_workbook(
            BytesIO(body), read_only=True, data_only=True, keep_links=False
        )
        try:
            source = book[sheet]
            if (
                source.max_row is None
                or source.max_row > MAX_SHEET_ROWS
                or source.max_column is None
                or source.max_column > MAX_SHEET_COLUMNS
            ):
                msg = "worksheet exceeds its row or column bound"
                raise ValueError(msg)
            return tuple(source.iter_rows(values_only=True))
        finally:
            book.close()

    return call_scrubbed(
        read, lambda message: ResponseError(message, provider=provider)
    )


def legacy_spreadsheet(
    body: bytes, sheet: str, provider: str
) -> tuple[tuple[Any, ...], ...]:
    """Read XLS in memory; macros and external links are ignored."""

    def read() -> tuple[tuple[Any, ...], ...]:
        import xlrd  # noqa: PLC0415 - lazy legacy file reader

        if len(body) > MAX_EXPANDED_BYTES:
            msg = "legacy workbook exceeds the size cap"
            raise ValueError(msg)
        book = xlrd.open_workbook(file_contents=body, on_demand=True)
        try:
            source = book.sheet_by_name(sheet)
            if source.nrows > MAX_SHEET_ROWS or source.ncols > MAX_SHEET_COLUMNS:
                msg = "legacy worksheet exceeds its row or column bound"
                raise ValueError(msg)
            return tuple(tuple(source.row_values(i)) for i in range(source.nrows))
        finally:
            book.release_resources()

    return call_scrubbed(
        read, lambda message: ResponseError(message, provider=provider)
    )
