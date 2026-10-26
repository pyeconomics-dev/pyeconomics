# tests/parity/test_parity_exports.py
# Copyright 2026 Nathan Ramos, CFA
# SPDX-License-Identifier: Apache-2.0
"""Every export reads back to the Python API's values (Phase 4 Step 10).

For one golden case per domain and one bound model, each export format (CSV,
XLSX, Parquet, Arrow IPC and JSON) is written by the Python API
(``pyeconomics.exports.export_bytes``), the command line (``models run
--format``) and the REST API (``?format=``), and read back as a user would read
it (pandas, openpyxl, pyarrow):

- every value equals the Python API's exactly: a number to the last bit, and
  text, a date or an array as the format spells it (ADR-0012 decision 6: the
  formula escape and canonical JSON for CSV and XLSX, lists for Parquet and
  Arrow);
- the manifest read back equals the result's: the CSV's ``#`` block, the XLSX
  manifest sheet, the Parquet and Arrow schema metadata, the JSON document;
- the CSV, Parquet, Arrow and JSON bytes are the same from every surface; an
  XLSX file records the second it was written, so it is compared by its sheets.

The bound model's exports, like its runs, carry no canary key.
"""

from __future__ import annotations

import io
import json
import math
from typing import TYPE_CHECKING, Any, TypeGuard

import pytest
from _loader import audit
from export_support import arrow_metadata, arrow_table, csv_metadata, read_csv
from extra_gaps import excluded_gap
from leakcheck import LeakCheck, canary
from openpyxl import load_workbook
from parity_support import CliSession, Run
from parity_world import AS_OF, inputs_for, parity_world
from server_support import LOOPBACK_ONLY, client

import pyeconomics
from pyeconomics import __version__, registry
from pyeconomics.core import tables
from pyeconomics.core.binding import DEFAULT_PROFILE
from pyeconomics.core.canonical import canonical_json
from pyeconomics.core.provenance import LicenceClass
from pyeconomics.core.registry import installed
from pyeconomics.data import cache
from pyeconomics.exports import export_bytes
from pyeconomics.exports.layout import cell_text, escape_text

if TYPE_CHECKING:
    from collections.abc import Iterator, Mapping

    from config_isolation import Isolation
    from fastapi.testclient import TestClient

    from pyeconomics.core.results import Result

pytestmark = LOOPBACK_ONLY

FORMATS = ("csv", "xlsx", "parquet", "arrow", "json")
#: Formats whose bytes are the same whoever writes them.
STABLE = ("csv", "parquet", "arrow", "json")
#: The labels of the texts the data must be shown with.
PROVENANCE = ("attribution", "notice", "terms")
#: The bound model whose exports are read back.
BOUND = ("macro.policy_rule", "taylor_1993")


def _domain_cases() -> list[Run]:
    """The first golden case of each domain's first model that runs here."""
    chosen: dict[str, Run] = {}
    for golden in audit(installed()).files:
        model = registry.get(golden.model)
        domain = model.spec.domain
        if domain not in chosen and excluded_gap(model) is None:
            chosen[domain] = Run(golden.model, golden.cases[0].inputs)
    return list(chosen.values())


DOMAIN_RUNS = _domain_cases()
DOMAIN_CASES = [pytest.param(run, id=run.model_id) for run in DOMAIN_RUNS]


@pytest.fixture(scope="module")
def cli() -> CliSession:
    return CliSession()


@pytest.fixture(scope="module")
def api() -> Iterator[TestClient]:
    with client() as test_client:
        yield test_client


# --- writing every export on every surface -------------------------------------------


def _options(run: Run) -> list[str]:
    if not run.bound:
        return []
    assert run.as_of is not None
    return ["--as-of", run.as_of.isoformat(), "--profile", run.profile or ""]


def _exports(
    run: Run, result: Result[Any, Any], cli: CliSession, api: TestClient
) -> dict[str, dict[str, bytes]]:
    """Each format's bytes from the Python API, the command line and the REST API."""
    params = (
        {"as_of": AS_OF.isoformat(), "profile": DEFAULT_PROFILE} if run.bound else {}
    )
    written: dict[str, dict[str, bytes]] = {}
    for export_format in FORMATS:
        cache.clear()
        outcome = cli(
            "models",
            "run",
            run.model_id,
            "--inputs",
            "-",
            "--format",
            export_format,
            *_options(run),
            stdin=run.body,
        )
        assert (outcome.code, outcome.stderr) == (0, "")
        cache.clear()
        response = api.post(
            f"/v1/models/{run.model_id}/run",
            content=run.body,
            params={**params, "format": export_format},
        )
        assert response.status_code == 200, response.text
        printed = outcome.stdout
        if export_format == "json":  # the canonical bytes, then one newline
            assert printed.endswith(b"\n")
            printed = printed[:-1]
        written[export_format] = {
            "python": export_bytes(result, export_format),
            "cli": printed,
            "rest": response.content,
        }
    return written


# --- what each format holds, read back -----------------------------------------------


def _spelled(value: object) -> object:
    """A value as CSV and XLSX spell it: numbers as they are, else escaped text."""
    if value is None or isinstance(value, bool | int | float):
        return value
    return escape_text(cell_text(value))


def _same_number(found: object, expected: float) -> bool:
    return (
        isinstance(found, int | float)
        and not isinstance(found, bool)
        and not math.isnan(found)
        and found == expected
        and math.copysign(1.0, found) == math.copysign(1.0, expected)
    )


def _same_cell(found: object, expected: object) -> bool:
    """Whether a read-back cell is a value as CSV or XLSX spells it."""
    spelled = _spelled(expected)
    if spelled is None:
        return found is None or (isinstance(found, float) and math.isnan(found))
    if isinstance(spelled, bool):
        return bool(found) is spelled and str(found).lower() == str(spelled).lower()
    if isinstance(spelled, int | float):
        return _same_number(found, float(spelled))
    return found == spelled


def _values(result: Result[Any, Any]) -> dict[str, object]:
    """The Python API's values, in the exports' columns."""
    return {name: values[0] for name, values in tables.columns_of([result]).items()}


def _check_csv(data: bytes, result: Result[Any, Any]) -> None:
    frame = read_csv(
        data, float_precision="round_trip", keep_default_na=False, na_values=[""]
    )
    values = _values(result)
    assert list(frame.columns) == list(values)
    row = frame.iloc[0].to_dict()
    differing = [name for name, v in values.items() if not _same_cell(row[name], v)]
    assert differing == []
    manifest = result.manifest
    assert csv_metadata(data) == [
        ("pyeconomics", manifest.package_version),
        ("model", f"{result.model_id} version {result.model_version}"),
        ("inputs_sha256", manifest.inputs_sha256),
        ("result_sha256", manifest.result_sha256),
        *_provenance(result),
    ]


def _provenance(result: Result[Any, Any]) -> list[tuple[str, str]]:
    """The attribution, notice and terms lines a result's data must be shown with."""
    kept = [
        record
        for record in result.manifest.data_sources
        if record.licence is not LicenceClass.USER_SUPPLIED
    ]
    return [
        *(("attribution", text) for text in dict.fromkeys(r.attribution for r in kept)),
        *(
            ("notice", text)
            for text in dict.fromkeys(r.notice for r in kept if r.notice)
        ),
        *(
            ("terms", url)
            for url in dict.fromkeys(r.terms_url for r in kept if r.terms_url)
        ),
    ]


def _rows(data: bytes, sheet: str) -> list[list[object]]:
    """A sheet's rows of values, each without the empty cells that end it."""
    rows = []
    for row in load_workbook(io.BytesIO(data))[sheet].iter_rows(values_only=True):
        values: list[object] = list(row)
        while values and values[-1] is None:
            values.pop()
        rows.append(values)
    return rows


def _flatten(
    mapping: Mapping[str, object], prefix: str = ""
) -> Iterator[tuple[str, object]]:
    for key, value in mapping.items():
        if isinstance(value, dict):
            yield from _flatten(value, f"{prefix}{key}.")
        else:
            yield f"{prefix}{key}", value


def _is_records(value: object) -> TypeGuard[list[dict[str, object]]]:
    return (
        isinstance(value, list)
        and bool(value)
        and all(isinstance(item, dict) for item in value)
    )


def _read_manifest_sheet(
    rows: list[list[object]],
) -> tuple[list[list[object]], dict[str, object], dict[str, list[dict[str, object]]]]:
    """Read the manifest sheet back: provenance, fields, and record tables.

    The sheet holds the provenance lines, a blank row, a ``[name, value]`` row
    per field, then each list of records as a blank row, its name, its keys and
    one row per record.
    """
    provenance = []
    while rows and rows[0] and rows[0][0] in PROVENANCE:
        provenance.append(rows.pop(0))
    if provenance and rows and rows[0] == []:
        rows.pop(0)
    fields: dict[str, object] = {}
    while rows and rows[0] != []:
        name, *value = rows.pop(0)
        fields[str(name)] = value[0] if value else None
    records: dict[str, list[dict[str, object]]] = {}
    while rows:
        assert rows.pop(0) == []
        (name,), keys = rows.pop(0), [str(key) for key in rows.pop(0)]
        found = []
        while rows and rows[0] != []:
            cells = rows.pop(0)
            found.append(
                {
                    key: cells[i] if i < len(cells) else None
                    for i, key in enumerate(keys)
                }
            )
        records[str(name)] = found
    return provenance, fields, records


def _check_xlsx(data: bytes, result: Result[Any, Any]) -> None:
    header, row = _rows(data, "data")
    values = _values(result)
    assert header == list(values)
    row += [None] * (len(header) - len(row))
    differing = [
        name
        for name, cell in zip(header, row, strict=True)
        if not _same_cell(cell, values[str(name)])
    ]
    assert differing == []
    provenance, fields, records = _read_manifest_sheet(_rows(data, "manifest"))
    assert [tuple(line) for line in provenance] == _provenance(result)
    manifest = json.loads(canonical_json(result.manifest.to_dict()))
    flat = dict(_flatten(manifest))
    assert fields == {
        name: _spelled(value) for name, value in flat.items() if not _is_records(value)
    }
    assert records == {
        name: [{key: _spelled(v) for key, v in record.items()} for record in value]
        for name, value in flat.items()
        if _is_records(value)
    }


def _check_arrow(data: bytes, export_format: str, result: Result[Any, Any]) -> None:
    (row,) = arrow_table(data, export_format).to_pylist()
    values = _values(result)
    assert list(row) == list(values)
    assert row == values
    metadata = arrow_metadata(data, export_format)
    assert metadata["manifest"] == json.loads(canonical_json(result.manifest.to_dict()))
    assert metadata["attributions"] == list(result.attributions())
    assert metadata["version"] == __version__


def _check(export_format: str, data: bytes, result: Result[Any, Any]) -> None:
    if export_format == "csv":
        _check_csv(data, result)
    elif export_format == "xlsx":
        _check_xlsx(data, result)
    elif export_format == "json":
        assert data == result.to_json().encode("utf-8")
    else:
        _check_arrow(data, export_format, result)


def _check_all(
    run: Run, result: Result[Any, Any], cli: CliSession, api: TestClient
) -> dict[str, dict[str, bytes]]:
    written = _exports(run, result, cli, api)
    for export_format, by_surface in written.items():
        for data in by_surface.values():
            _check(export_format, data, result)
        if export_format in STABLE:
            assert by_surface["cli"] == by_surface["python"], export_format
            assert by_surface["rest"] == by_surface["python"], export_format
    return written


@pytest.mark.parametrize("run", DOMAIN_CASES)
def test_a_domains_exports_read_back_to_its_values(
    run: Run, cli: CliSession, api: TestClient
) -> None:
    result = pyeconomics.run(run.model_id, dict(run.inputs))
    _check_all(run, result, cli, api)


def test_one_case_per_domain() -> None:
    domains = {model.spec.domain for model in registry.models()}
    assert {registry.get(run.model_id).spec.domain for run in DOMAIN_RUNS} == domains


def test_a_bound_models_exports_read_back_to_its_values_and_records(
    isolated: Isolation,
    monkeypatch: pytest.MonkeyPatch,
    cli: CliSession,
    api: TestClient,
) -> None:
    world = parity_world(monkeypatch)
    keys = []
    for info in world.infos():
        if info.key_variable is not None:
            keys.append(canary(info.id))
            isolated.keyring.set_password("pyeconomics", info.id, keys[-1].value)
    run = Run(BOUND[0], inputs_for(BOUND), as_of=AS_OF, profile=DEFAULT_PROFILE)
    cache.clear()
    overrides: dict[str, Any] = dict(run.inputs)
    bound = pyeconomics.bind(run.model_id, as_of=AS_OF, **overrides)
    result = pyeconomics.run(run.model_id, bound)
    assert len(_provenance(result)) > 2  # every source's attribution and terms
    written = _check_all(run, result, cli, api)
    for planted in keys:
        check = LeakCheck(planted)
        for export_format, by_surface in written.items():
            for surface, data in by_surface.items():
                check.add(f"{surface}:{export_format}", data)
        check.assert_clean()
