# tests/exports/export_support.py
# Copyright 2026 Nathan Ramos, CFA
# SPDX-License-Identifier: Apache-2.0
"""Read exports back the way a user would (tests/exports, Phase 4 Step 5).

Each reader takes an export's bytes: :func:`read_csv` through
``pandas.read_csv(..., comment="#")`` and :func:`csv_metadata` for its ``#``
block; :func:`xlsx_sheet` and :func:`xlsx_cells` through openpyxl;
:func:`arrow_metadata` through pyarrow, for Parquet and Arrow IPC alike.
:data:`HOSTILE` holds text cells a spreadsheet would read as formulas, or that
break a careless CSV writer.
"""

from __future__ import annotations

import importlib
import io
import json
from typing import Any, Final

import pandas as pd
from openpyxl import load_workbook

__all__ = [
    "HOSTILE",
    "NUMBERS",
    "arrow_metadata",
    "arrow_table",
    "csv_metadata",
    "read_csv",
    "xlsx_cells",
    "xlsx_sheet",
]

#: Text cells: each formula prefix, then the characters CSV must quote.
HOSTILE: Final = (
    "=1+1",
    "+1",
    "-1",
    "@SUM(A1)",
    "\tTAB",
    "\rCR",
    "#hash",
    "a,b",
    'say "hi"',
    "two\nlines",
    "plain",
)
#: Numbers a spreadsheet must keep as numbers: negatives and exponents.
NUMBERS: Final = (-1.5, -3, 1e-07, -2.5e20, 0.1, 7, -0.0, 6.02e23, 1e16, -42, 3.25)


def read_csv(data: bytes, **options: Any) -> pd.DataFrame:  # noqa: ANN401
    """Read a CSV export as a user would: skipping its ``#`` block.

    pandas' default float parser can miss a value's last bit; pass
    ``float_precision="round_trip"`` to read every number back exactly.
    """
    frame: pd.DataFrame = pd.read_csv(io.BytesIO(data), comment="#", **options)
    return frame


def csv_metadata(data: bytes) -> list[tuple[str, str]]:
    """Return the ``# label: text`` lines of a CSV export, in order."""
    found = []
    for line in data.decode("utf-8").split("\n"):
        if not line.startswith("# "):
            break
        label, _, text = line[2:].partition(": ")
        found.append((label, text))
    return found


def xlsx_sheet(data: bytes, name: str) -> list[list[Any]]:
    """Return a sheet's rows of values, without the empty cells that end a row.

    openpyxl pads every row to the sheet's widest.
    """
    workbook = load_workbook(io.BytesIO(data))
    rows = []
    for row in workbook[name].iter_rows(values_only=True):
        values = list(row)
        while values and values[-1] is None:
            values.pop()
        rows.append(values)
    return rows


def xlsx_cells(data: bytes) -> list[Any]:
    """Return every cell of every sheet, as openpyxl reads it."""
    workbook = load_workbook(io.BytesIO(data))
    return [cell for sheet in workbook for row in sheet.iter_rows() for cell in row]


def arrow_table(data: bytes, export_format: str) -> Any:  # noqa: ANN401
    """Read a Parquet or Arrow IPC export into an Arrow table.

    pyarrow is imported by name: it ships no type information, and the
    test group installs it on every cell (a missing one fails the test).
    """
    if export_format == "parquet":
        return importlib.import_module("pyarrow.parquet").read_table(io.BytesIO(data))
    ipc = importlib.import_module("pyarrow.ipc")
    return ipc.open_file(io.BytesIO(data)).read_all()


def arrow_metadata(data: bytes, export_format: str) -> dict[str, Any]:
    """Return a Parquet or Arrow export's pyeconomics metadata, decoded."""
    metadata = arrow_table(data, export_format).schema.metadata
    return {
        "manifest": json.loads(metadata[b"pyeconomics.manifest"]),
        "attributions": json.loads(metadata[b"pyeconomics.attributions"]),
        "version": metadata[b"pyeconomics.version"].decode("utf-8"),
    }
